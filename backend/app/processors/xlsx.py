"""
XLSX document processor.

The processor extracts workbook structure into the common
DocumentContent / DocumentSegment representation.

Responsibilities:
    - Open XLSX workbooks
    - Extract worksheets
    - Extract non-empty rows
    - Preserve cell coordinates
    - Preserve column information
    - Preserve raw cell values
    - Preserve Excel data types
    - Preserve formulas

Non-responsibilities:
    - Financial interpretation
    - Financial calculations
    - Evidence generation
    - Investment dimension classification
    - Source reconciliation
    - LLM invocation
"""

from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID

import openpyxl
from openpyxl.cell.cell import TYPE_BOOL, TYPE_ERROR, TYPE_FORMULA

from app.processors.base import (
    DocumentContent,
    DocumentProcessor,
    DocumentSegment,
)


class XlsxProcessor(DocumentProcessor):
    """
    Processor for Microsoft Excel XLSX workbooks.
    """

    _SUPPORTED_MIME_TYPES = {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }

    _SUPPORTED_EXTENSIONS = {
        ".xlsx",
    }

    @property
    def supported_mime_types(self) -> set[str]:
        return self._SUPPORTED_MIME_TYPES

    @property
    def supported_extensions(self) -> set[str]:
        return self._SUPPORTED_EXTENSIONS

    def process(
        self,
        document_id: UUID,
        path: Path,
    ) -> DocumentContent:
        """
        Extract workbook content while preserving worksheet/cell structure.

        Each non-empty worksheet row becomes one DocumentSegment.

        The segment text remains a deterministic human-readable
        representation of the row.

        Structural information is stored in segment.metadata["cells"].
        """

        workbook = openpyxl.load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
        )

        try:
            segments: list[DocumentSegment] = []
            text_parts: list[str] = []

            workbook_metadata: dict[str, Any] = {
                "filename": path.name,
                "extension": path.suffix.lower(),
                "format": "xlsx",
                "sheet_count": len(workbook.sheetnames),
                "sheet_names": list(workbook.sheetnames),
                "formula_preserved": True,
                "data_only": False,
                "sheets": [],
            }

            segment_index = 0

            for sheet_index, sheet_name in enumerate(
                workbook.sheetnames
            ):
                worksheet = workbook[sheet_name]

                sheet_metadata = {
                    "sheet_index": sheet_index,
                    "sheet_name": sheet_name,
                    "max_row": 0,
                    "max_column": 0,
                    "nonempty_rows": 0,
                    "nonempty_cells": 0,
                    "formula_cells": 0,
                }

                for row_number, row in enumerate(
                    worksheet.iter_rows(),
                    start=1,
                ):
                    cells = self._extract_cells(row)

                    if not cells:
                        continue

                    sheet_metadata["max_row"] = max(
                        sheet_metadata["max_row"],
                        row_number,
                    )

                    sheet_metadata["max_column"] = max(
                        sheet_metadata["max_column"],
                        max(
                            cell["column_index"]
                            for cell in cells
                        ),
                    )

                    sheet_metadata["nonempty_rows"] += 1
                    sheet_metadata["nonempty_cells"] += len(cells)

                    sheet_metadata["formula_cells"] += sum(
                        1
                        for cell in cells
                        if cell["data_type"] == "f"
                    )

                    row_text = self._format_row(cells)

                    if not row_text:
                        continue

                    if text_parts:
                        text_parts.append("\n\n")

                    text_parts.append(
                        f"[Sheet: {sheet_name} | Row: {row_number}]\n"
                        f"{row_text}"
                    )

                    segments.append(
                        DocumentSegment(
                            index=segment_index,
                            text=row_text,
                            start_offset=0,
                            end_offset=len(row_text),
                            metadata={
                                "type": "worksheet_row",
                                "sheet_index": sheet_index,
                                "sheet_name": sheet_name,
                                "row_number": row_number,
                                "cell_count": len(cells),
                                "cells": cells,
                            },
                        )
                    )

                    segment_index += 1

                workbook_metadata["sheets"].append(
                    sheet_metadata
                )

            text = "".join(text_parts)

            workbook_metadata["segment_count"] = len(segments)
            workbook_metadata["text_length"] = len(text)
            workbook_metadata["formula_cell_count"] = sum(
                sheet["formula_cells"]
                for sheet in workbook_metadata["sheets"]
            )

            return DocumentContent(
                document_id=document_id,
                title=path.stem,
                text=text,
                page_count=len(workbook.sheetnames),
                metadata=workbook_metadata,
                segments=tuple(segments),
            )

        finally:
            workbook.close()

    @classmethod
    def _extract_cells(
        cls,
        row,
    ) -> list[dict[str, Any]]:
        """
        Extract non-empty cells from one worksheet row.

        Cell position and raw value are preserved.

        Important:
            Empty cells are NOT included in the structural cell list.
            Their absence is intentional and should not be interpreted
            as a zero or empty financial value.
        """

        cells: list[dict[str, Any]] = []

        for cell in row:
            if cell.value is None:
                continue

            value = cls._normalize_cell_value(cell.value)

            if value == "":
                continue

            cells.append(
                {
                    "coordinate": cell.coordinate,
                    "column": cell.column_letter,
                    "column_index": cell.column,
                    "row": cell.row,
                    "value": value,
                    "data_type": cls._normalize_data_type(
                        cell.data_type
                    ),
                    "is_formula": cell.data_type == TYPE_FORMULA,
                }
            )

        return cells

    @staticmethod
    def _normalize_data_type(
        data_type: str | None,
    ) -> str:
        """
        Normalize openpyxl cell data types.

        Common values:
            s = string
            n = number
            f = formula
            b = boolean
            e = error
        """

        if data_type == TYPE_FORMULA:
            return "f"

        if data_type == TYPE_BOOL:
            return "b"

        if data_type == TYPE_ERROR:
            return "e"

        if data_type is None:
            return "unknown"

        return str(data_type)

    @staticmethod
    def _normalize_cell_value(
        value: Any,
    ) -> str:
        """
        Convert an XLSX cell value to deterministic text.

        This function does not interpret financial meaning.
        """

        if value is None:
            return ""

        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"

        if isinstance(value, datetime):
            return value.isoformat(sep=" ")

        if isinstance(value, date):
            return value.isoformat()

        if isinstance(value, time):
            return value.isoformat()

        if isinstance(value, Decimal):
            return str(value)

        if isinstance(value, float):
            return str(value)

        return str(value).strip()

    @staticmethod
    def _format_row(
        cells: list[dict[str, Any]],
    ) -> str:
        """Produce the human-readable row representation."""

        if not cells:
            return ""

        first_column = min(
            cell["column_index"]
            for cell in cells
        )

        last_column = max(
            cell["column_index"]
            for cell in cells
        )

        values_by_column = {
            cell["column_index"]: cell["value"]
            for cell in cells
        }

        values = [
            values_by_column.get(column_index, "")
            for column_index in range(
                first_column,
                last_column + 1
            )
        ]

        return " | ".join(values).strip()
